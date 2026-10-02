"""
Name: excel_to_syllabus.py
Purpose: Import syllabuses, modules and events from a flat Excel sheet (one row per event).
Created: 2026-10-02
Author: Michael K. Steinberg

Needs `pip install "bluz[excel]"` (openpyxl).

The sheet has a header row; column order does not matter and headers may be
English or Hebrew:

    | Syllabus   | Module     | Event              | Type   | Minutes | Shuffles | Critical |
    | סילבוס      | מערך        | מופע                | סוג     | דקות     | שאפלים    | קריטי     |
    |------------|------------|--------------------|--------|---------|----------|----------|
    | Networking | Foundations| OSI model          | הרצאה  | 90      |          |          |
    | Networking | Foundations| Packet capture lab | ע"ע    | 120     | A,B      | yes      |

Rows with an empty Event cell only make sure the syllabus/module exists.
Re-running is safe: syllabuses, modules and events already present (same
title, same parent) are reused, not duplicated.

Usage:
    python -m bluz.examples excel_to_syllabus              # writes a template + dry run
    BLUZ_EXCEL=plan.xlsx python -m bluz.examples excel_to_syllabus
"""

from __future__ import annotations

import os
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from bluz import Bluz, Curriculum, ModuleEventType

# Header spellings → canonical column names.
HEADERS: dict[str, str] = {
    "syllabus": "syllabus",
    "סילבוס": "syllabus",
    "מקצוע": "syllabus",
    "module": "module",
    "מערך": "module",
    "event": "event",
    "מופע": "event",
    "שיעור": "event",
    "type": "type",
    "סוג": "type",
    "minutes": "minutes",
    "דקות": "minutes",
    "duration": "minutes",
    "משך": "minutes",
    "shuffles": "shuffles",
    "שאפלים": "shuffles",
    "critical": "critical",
    "קריטי": "critical",
    "comment": "comment",
    "הערה": "comment",
}

# Event-type spellings (enum value, enum name, English word) → enum member.
TYPES: dict[str, ModuleEventType] = {
    **{member.value: member for member in ModuleEventType},
    **{member.name.lower(): member for member in ModuleEventType},
    "exercise": ModuleEventType.EXERCISE,
    "self study": ModuleEventType.SELF_TEACHING,
}

TRUE_WORDS = {"1", "yes", "y", "true", "x", "כן", "v", "✓"}


@dataclass
class EventRow:
    """One event of the outline."""

    title: str
    type: ModuleEventType = ModuleEventType.LECTURE
    minutes: int = 0
    shuffles: list[str] = field(default_factory=list)
    critical: bool = False
    comment: str | None = None


@dataclass
class Outline:
    """syllabus title → module title → events, in sheet order."""

    syllabuses: dict[str, dict[str, list[EventRow]]] = field(default_factory=dict)

    def add(self, syllabus: str, module: str | None, event: EventRow | None) -> None:
        modules = self.syllabuses.setdefault(syllabus, {})
        if module:
            events = modules.setdefault(module, [])
            if event:
                events.append(event)

    @property
    def event_count(self) -> int:
        return sum(len(e) for m in self.syllabuses.values() for e in m.values())


class SheetError(ValueError):
    """A row the importer cannot understand — carries the sheet row number."""


def parse_type(value: Any, *, row: int) -> ModuleEventType:
    if value in (None, ""):
        return ModuleEventType.LECTURE
    key = str(value).strip()
    found = TYPES.get(key) or TYPES.get(key.lower())
    if found is None:
        raise SheetError(
            f"row {row}: unknown event type {key!r} (use one of {sorted(TYPES)})"
        )
    return found


def parse_minutes(value: Any, *, row: int) -> int:
    if value in (None, ""):
        return 0
    try:
        return int(float(value))
    except (TypeError, ValueError):
        raise SheetError(
            f"row {row}: minutes must be a number, got {value!r}"
        ) from None


def rows_to_outline(rows: Iterable[tuple[Any, ...]]) -> Outline:
    """Turn sheet rows (header first) into an `Outline`. Pure — no Excel needed,
    so it is easy to test and to reuse with CSV or Google Sheets data."""
    iterator: Iterator[tuple[Any, ...]] = iter(rows)
    header = next(iterator, None)
    if header is None:
        raise SheetError("the sheet is empty")
    columns = {
        HEADERS[str(cell).strip().lower()]: index
        for index, cell in enumerate(header)
        if cell is not None and str(cell).strip().lower() in HEADERS
    }
    missing = {"syllabus", "module"} - set(columns)
    if missing:
        raise SheetError(f"missing column(s): {', '.join(sorted(missing))}")

    def cell(values: tuple[Any, ...], name: str) -> Any:
        index = columns.get(name)
        return values[index] if index is not None and index < len(values) else None

    outline = Outline()
    for number, values in enumerate(iterator, start=2):
        syllabus = str(cell(values, "syllabus") or "").strip()
        if not syllabus:
            continue  # blank spacer row
        module = str(cell(values, "module") or "").strip() or None
        title = str(cell(values, "event") or "").strip()
        event = None
        if title:
            shuffles = str(cell(values, "shuffles") or "")
            event = EventRow(
                title=title,
                type=parse_type(cell(values, "type"), row=number),
                minutes=parse_minutes(cell(values, "minutes"), row=number),
                shuffles=[
                    s.strip()
                    for s in shuffles.replace(";", ",").split(",")
                    if s.strip()
                ],
                critical=str(cell(values, "critical") or "").strip().lower()
                in TRUE_WORDS,
                comment=(str(cell(values, "comment")).strip() or None)
                if cell(values, "comment")
                else None,
            )
        outline.add(syllabus, module, event)
    return outline


def read_outline(path: str | Path, sheet: str | None = None) -> Outline:
    """Read an `Outline` from an .xlsx file (first sheet unless named)."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet] if sheet else workbook.worksheets[0]
        return rows_to_outline(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()


def write_template(path: str | Path) -> Path:
    """Write an example workbook in the expected layout."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Syllabuses"
    sheet.append(
        ["Syllabus", "Module", "Event", "Type", "Minutes", "Shuffles", "Critical"]
    )
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.append(["Networking", "Foundations", "OSI model", "הרצאה", 90, "", ""])
    sheet.append(
        ["Networking", "Foundations", "Packet capture lab", 'ע"ע', 120, "A,B", "yes"]
    )
    sheet.append(
        ["Networking", "Routing", "Static & dynamic routing", "lecture", 90, "", ""]
    )
    sheet.append(["Networking", "Routing", "BGP reading", "self_teaching", 60, "", ""])
    sheet.append(["Security", "Basics", "", "", "", "", ""])
    target = Path(path)
    workbook.save(target)
    return target


@dataclass
class ImportReport:
    created: dict[str, int] = field(
        default_factory=lambda: {"syllabus": 0, "module": 0, "event": 0}
    )
    reused: dict[str, int] = field(
        default_factory=lambda: {"syllabus": 0, "module": 0, "event": 0}
    )

    def __str__(self) -> str:
        return ", ".join(
            f"{kind}: {self.created[kind]} new / {self.reused[kind]} existing"
            for kind in self.created
        )


def import_outline(
    curriculum: Curriculum, outline: Outline, *, dry_run: bool = False
) -> ImportReport:
    """Create what `outline` describes under `curriculum`, reusing what exists.

    With `dry_run=True` nothing is written; the report says what would be.
    """
    report = ImportReport()
    for syllabus_title, modules in outline.syllabuses.items():
        syllabus = curriculum.syllabuses.find(syllabus_title)
        if syllabus is not None:
            report.reused["syllabus"] += 1
        else:
            report.created["syllabus"] += 1
            if not dry_run:
                syllabus = curriculum.create_syllabus(syllabus_title)
        for module_title, events in modules.items():
            module = (
                syllabus.modules.find(module_title)
                if syllabus is not None and syllabus.links is not None
                else None
            )
            if module is not None:
                report.reused["module"] += 1
            else:
                report.created["module"] += 1
                if not dry_run and syllabus is not None:
                    module = syllabus.create_module(module_title)
            existing = (
                {e.title for e in module.events}
                if module is not None and module.links is not None
                else set()
            )
            for row in events:
                if row.title in existing:
                    report.reused["event"] += 1
                    continue
                report.created["event"] += 1
                if not dry_run and module is not None:
                    module.create_event(
                        row.title,
                        type=row.type,
                        minimum_duration=row.minutes,
                        shuffles=row.shuffles or None,
                        is_critical=row.critical or None,
                        comment=row.comment,
                    )
    return report


def main(bz: Bluz, path: str | None = None, *, dry_run: bool = True) -> None:
    source = path or os.getenv("BLUZ_EXCEL")
    if source is None:
        source = str(write_template("bluz-syllabus-template.xlsx"))
        print(f"No BLUZ_EXCEL given — wrote a template to {source}")
    outline = read_outline(source)
    print(
        f"{len(outline.syllabuses)} syllabuses, {outline.event_count} events in {source}"
    )

    curriculum = bz.iterations.current().curriculum
    if curriculum is None:
        print("The current iteration has no linked curriculum.")
        return
    report = import_outline(curriculum, outline, dry_run=dry_run)
    print(
        ("Would import" if dry_run else "Imported")
        + f" into {curriculum.title!r}: {report}"
    )
    if dry_run:
        print("Run main(bz, path, dry_run=False) to write.")


if __name__ == "__main__":
    with Bluz() as session:
        main(session)
