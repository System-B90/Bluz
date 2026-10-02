"""
Name: excel_wide_syllabus.py
Purpose: Import a hand-made workbook — one sheet per syllabus, bold module header rows.
Created: 2026-10-02
Author: Michael K. Steinberg

Needs `pip install "bluz[excel]"` (openpyxl).

Many syllabuses arrive as human-formatted workbooks rather than flat tables:

    Sheet "Networking"
    | Foundations          |        |     |     <- module row: only column A, bold
    | OSI model            | הרצאה  | 90  |     <- event rows: title, type, minutes
    | Packet capture lab   | ע"ע    | 120 |
    |                      |        |     |     <- blank rows are ignored
    | Routing              |        |     |
    | Static routing       | הרצאה  | 90  |

This example turns that into the same `Outline` the flat importer uses, so
`import_outline()` (create-or-reuse, dry run) is shared rather than repeated.
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any

from bluz import Bluz
from bluz.examples.excel_to_syllabus import (
    EventRow,
    Outline,
    import_outline,
    parse_minutes,
    parse_type,
)


def _is_module_row(cells: list[Any]) -> bool:
    """A module header: text in column A and nothing else on the row.

    Bold is how people *draw* headers, but it is not reliable (copied sheets
    lose it), so the shape of the row decides.
    """
    rest = [c for c in cells[1:] if c not in (None, "")]
    return bool(cells and cells[0]) and not rest


def read_wide_outline(path: str | Path) -> Outline:
    """One sheet per syllabus; module header rows followed by event rows."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    outline = Outline()
    for sheet in workbook.worksheets:
        module: str | None = None
        outline.add(sheet.title, None, None)
        for row in sheet.iter_rows():
            cells = [cell.value for cell in row]
            if not any(c not in (None, "") for c in cells):
                continue
            if _is_module_row(cells):
                module = str(cells[0]).strip()
                outline.add(sheet.title, module, None)
                continue
            if module is None:
                continue  # rows above the first module header (titles, notes)
            number = getattr(row[0], "row", 0)
            title = str(cells[0]).strip()
            outline.add(
                sheet.title,
                module,
                EventRow(
                    title=title,
                    type=parse_type(cells[1] if len(cells) > 1 else None, row=number),
                    minutes=parse_minutes(
                        cells[2] if len(cells) > 2 else None, row=number
                    ),
                ),
            )
    workbook.close()
    return outline


def write_wide_template(path: str | Path) -> Path:
    """Write an example workbook in the wide layout."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Networking"
    for values, bold in (
        (["Foundations"], True),
        (["OSI model", "הרצאה", 90], False),
        (["Packet capture lab", 'ע"ע', 120], False),
        ([], False),
        (["Routing"], True),
        (["Static routing", "הרצאה", 90], False),
    ):
        sheet.append(values)
        if bold:
            sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
    target = Path(path)
    workbook.save(target)
    return target


def main(bz: Bluz, path: str | None = None, *, dry_run: bool = True) -> None:
    source = path or os.getenv("BLUZ_EXCEL")
    if source is None:
        source = str(write_wide_template("bluz-wide-template.xlsx"))
        print(f"No BLUZ_EXCEL given — wrote a template to {source}")
    outline = read_wide_outline(source)
    for syllabus, modules in outline.syllabuses.items():
        print(
            f"{syllabus}: " + ", ".join(f"{m} ({len(e)})" for m, e in modules.items())
        )

    curriculum = bz.iterations.current().curriculum
    if curriculum is None:
        print("The current iteration has no linked curriculum.")
        return
    report = import_outline(curriculum, outline, dry_run=dry_run)
    print(("Would import: " if dry_run else "Imported: ") + str(report))


if __name__ == "__main__":
    with Bluz() as session:
        main(session)
